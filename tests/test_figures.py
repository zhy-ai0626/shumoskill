#!/usr/bin/env python3
"""`figures/cumcm_style.py` 的守卫测试，以及不依赖真题附件的 gallery 范例能出图。

这里只测**会静默出错的那些**，不测长相：

- `corr_heatmap` 的色标必须钉在 [-1, 1]。随手 `imshow(C)` 会按数据范围自动拉伸，
  0 就不在中点上了——图上"看起来中性"的格子其实是强正相关，**而且不报错**。
- `flow_arrow(angle=...)` 两端方向平行时没有交点，matplotlib 画出来是条乱线，
  也不报错。
- `finish()` 必须拦住空的轴标签（2022A 讲评点名"没有很好地呈现结果"）。
- `series_kw` 的槽位上限：第 9 个系列必须报错，而不是循环复用颜色。
"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "figures"))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt          # noqa: E402

import cumcm_style as cs                 # noqa: E402


def _corr(n: int, seed: int = 0) -> np.ndarray:
    rng = np.random.default_rng(seed)
    X = rng.normal(size=(200, n))
    return np.corrcoef(X, rowvar=False)


class CorrHeatmapGuardTest(unittest.TestCase):
    def setUp(self):
        cs.use()
        self.fig, self.ax = plt.subplots()

    def tearDown(self):
        plt.close(self.fig)

    def test_color_limits_are_pinned_to_pm_one(self):
        """**这是这个 helper 存在的唯一理由。** 相关系数都落在 0.3~0.9 时，
        自动拉伸会把色标中点放到 0.6，强正相关看起来像中性。"""
        M = np.array([[1.0, 0.8, 0.7], [0.8, 1.0, 0.6], [0.7, 0.6, 1.0]])
        im = cs.corr_heatmap(self.ax, M, ["a", "b", "c"])
        self.assertEqual(im.get_clim(), (-1.0, 1.0))

    def test_rejects_values_outside_pm_one(self):
        M = np.array([[4.0, 1.2], [1.2, 9.0]])          # 协方差，不是相关系数
        with self.assertRaises(ValueError) as ctx:
            cs.corr_heatmap(self.ax, M, ["a", "b"])
        self.assertIn("[-1, 1]", str(ctx.exception))

    def test_rejects_non_square(self):
        with self.assertRaises(ValueError):
            cs.corr_heatmap(self.ax, np.zeros((3, 4)), ["a", "b", "c"])

    def test_rejects_label_count_mismatch(self):
        with self.assertRaises(ValueError):
            cs.corr_heatmap(self.ax, _corr(4), ["a", "b"])

    def test_masks_upper_triangle_and_diagonal(self):
        M = _corr(5, seed=1)
        cs.corr_heatmap(self.ax, M, list("abcde"), mask_upper=True)
        shown = self.ax.images[0].get_array()
        for i in range(5):
            for j in range(5):
                if j >= i:
                    self.assertTrue(np.ma.is_masked(shown[i, j])
                                    or not np.isfinite(shown[i, j]),
                                    (i, j))

    def test_annotates_only_above_threshold(self):
        M = np.array([[1.0, 0.9, 0.1],
                      [0.9, 1.0, 0.2],
                      [0.1, 0.2, 1.0]])
        cs.corr_heatmap(self.ax, M, ["a", "b", "c"], annot_min=0.5,
                        cbar=False)
        texts = [t.get_text() for t in self.ax.texts]
        # 下三角里只有 0.9 过阈值；0.1 / 0.2 不该标
        self.assertEqual(texts, [".90"])

    def test_grid_is_off_on_major_axis(self):
        """rcParams 把 axes.grid 设成 True（数据图要网格）。
        热图上留着就是一层横条纹穿过所有格子。"""
        cs.corr_heatmap(self.ax, _corr(4, seed=2), list("abcd"))
        self.assertFalse(any(l.get_visible()
                             for l in self.ax.get_ygridlines()))


class FlowHelpersTest(unittest.TestCase):
    def setUp(self):
        cs.use()
        self.fig, self.ax = cs.flow_canvas()

    def tearDown(self):
        plt.close(self.fig)

    def test_canvas_has_no_axis_or_grid(self):
        self.assertFalse(self.ax.axison)

    def test_box_anchors(self):
        b = cs.flow_box(self.ax, 10, 20, 40, 10, "x")
        self.assertEqual(b.c, (30, 25))
        self.assertEqual(b.n, (30, 30))
        self.assertEqual(b.s, (30, 20))
        self.assertEqual(b.e, (50, 25))
        self.assertEqual(b.w_, (10, 25))
        self.assertEqual(b.top(0.25), (20, 30))
        self.assertEqual(b.bottom(0.75), (40, 20))

    def test_box_rejects_unknown_tint(self):
        with self.assertRaises(ValueError) as ctx:
            cs.flow_box(self.ax, 0, 0, 10, 5, "x", tint="rainbow")
        self.assertIn("tint", str(ctx.exception))

    def test_arrow_rejects_parallel_elbow(self):
        """两端方向平行 → 折线没有交点，画出来是条乱线，matplotlib 不报错。"""
        for bad in ((-90, 90), (0, 180), (90, 270)):
            with self.assertRaises(ValueError, msg=bad):
                cs.flow_arrow(self.ax, (0, 0), (10, 10), angle=bad)

    def test_arrow_accepts_perpendicular_elbow(self):
        cs.flow_arrow(self.ax, (0, 0), (10, 10), angle=(0, 90))
        self.assertGreaterEqual(len(self.ax.texts), 0)   # 没抛异常即可

    def test_arrow_label_offset_applied(self):
        cs.flow_arrow(self.ax, (0, 0), (6, 0), text="解沿用", text_dy=3.4)
        t = [t for t in self.ax.texts if t.get_text() == "解沿用"][0]
        self.assertAlmostEqual(t.get_position()[1], 3.4)


class FinishGuardTest(unittest.TestCase):
    def setUp(self):
        cs.use()
        self.fig, self.ax = plt.subplots()

    def tearDown(self):
        plt.close(self.fig)

    def test_requires_both_axis_labels(self):
        self.ax.plot([1, 2], [1, 2])
        for kw in ({"xlabel": "", "ylabel": "y (m)"},
                   {"xlabel": "x (s)", "ylabel": ""},
                   {}):
            with self.assertRaises(ValueError, msg=kw):
                cs.finish(self.ax, **kw)

    def test_ninth_series_raises_instead_of_reusing_color(self):
        for i in range(8):
            cs.series_kw(i, "line")
        with self.assertRaises(ValueError) as ctx:
            cs.series_kw(8, "line")
        self.assertIn("8 槽", str(ctx.exception))

    def test_scatter_series_cap(self):
        cs.series_kw(2, "scatter", n_series=3)
        with self.assertRaises(ValueError):
            cs.series_kw(3, "scatter", n_series=4)


class GalleryBuildTest(unittest.TestCase):
    """三张新范例都不依赖真题附件，任何人 clone 下来都该能出图。"""

    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, str(ROOT / "figures"))
        cls.prev_repo = os.environ.pop("CUMCM_REPO", None)
        import importlib
        import gallery
        cls.gallery = importlib.reload(gallery)

    @classmethod
    def tearDownClass(cls):
        if cls.prev_repo is not None:
            os.environ["CUMCM_REPO"] = cls.prev_repo

    def _build(self, key: str):
        g = self.gallery
        with tempfile.TemporaryDirectory() as d:
            prev, g.OUT = g.OUT, d
            try:
                cs.use()
                infos = g.BUILDERS[key]()
            finally:
                g.OUT = prev
            self.assertTrue(infos)
            for info in infos:
                self.assertTrue(os.path.getsize(info["path"]) > 5000,
                                info["path"])

    def test_flow(self):
        self._build("流程图")

    def test_corr(self):
        self._build("相关矩阵")

    def test_descriptive(self):
        self._build("描述统计")

    def test_attachment_free_builders_registered(self):
        for key in ("流程图", "相关矩阵", "描述统计"):
            self.assertIn(key, self.gallery.BUILDERS)


class FlowPptxTest(unittest.TestCase):
    """`flow_pptx.py` 的守卫。

    盯住的是 pptx 里"生成得出来、但其实没用"的那类问题：
    连接线没绑住两端（拖框不跟随——这是选 pptx 的唯一理由）、
    中文 run 没设 `<a:ea>`（中文静默变宋体）、`RGBColor` 是 tuple 子类
    导致 `"#%s" % rgb` 抛 TypeError（预览里每个框静默变白）。
    """

    @classmethod
    def setUpClass(cls):
        try:
            import pptx  # noqa: F401
        except ImportError:
            raise unittest.SkipTest("没装 python-pptx")
        sys.path.insert(0, str(ROOT / "figures"))
        import importlib
        import flow_pptx
        cls.mod = importlib.reload(flow_pptx)

    def _build(self, d: str) -> str:
        deck = self.mod.Deck()
        self.mod.slide_overall(deck)
        self.mod.slide_per_question(deck)
        return deck.save(os.path.join(d, "t.pptx"))

    def test_every_connector_is_bound_at_both_ends(self):
        with tempfile.TemporaryDirectory() as d:
            problems = self.mod.verify(self._build(d))
        self.assertEqual(problems, [], "\n".join(problems))

    def test_slide_size_is_figure_shaped_not_16x9(self):
        """幻灯片尺寸要按插图比例设，不是默认 16:9——
        否则"另存为图片"会带一大片空白，还得再裁一次。"""
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            prs = Presentation(self._build(d))
        self.assertAlmostEqual(prs.slide_width / 360000.0,
                               self.mod.SLIDE_W_CM, places=2)
        self.assertAlmostEqual(prs.slide_height / 360000.0,
                               self.mod.SLIDE_H_CM, places=2)

    def test_box_rejects_unknown_tint(self):
        deck = self.mod.Deck()
        deck.add_slide()
        with self.assertRaises(ValueError):
            deck.box(1, 1, 4, 1, "x", tint="rainbow")

    def test_tints_match_matplotlib_version(self):
        """pptx 版和 gallery 的 09 必须同配色，否则两版图放进同一篇论文会打架。"""
        from pptx import Presentation
        from pptx.enum.dml import MSO_FILL
        want = {v.lstrip("#").upper() for v in cs.FLOW_TINTS.values()}
        with tempfile.TemporaryDirectory() as d:
            prs = Presentation(self._build(d))
        got = {str(sh.fill.fore_color.rgb)
               for s in prs.slides for sh in s.shapes
               if not sh.element.tag.endswith("}cxnSp")
               and sh.fill.type == MSO_FILL.SOLID}
        self.assertTrue(got, "一个实心形状都没有——填充没设上")
        self.assertTrue(got <= want, "出现了 FLOW_TINTS 之外的颜色：%s" % (got - want))

    def test_preview_renders_every_solid_shape(self):
        """预览必须把每个实心形状的填充色都画出来。

        第一版用 `"#%s" % rgb` 取色：`RGBColor` 是 **tuple 子类**，
        `%` 会把三个分量当三个参数塞进一个占位符而抛 TypeError，
        再被 `except Exception` 兜成画布白——**预览里每个框都是白的，
        而 pptx 本身是对的**。这条盯住那次回归。
        """
        with tempfile.TemporaryDirectory() as d:
            pptx_path = self._build(d)
            png = os.path.join(d, "p.png")
            warns = self.mod.preview(pptx_path, png)
            self.assertEqual([w for w in warns if "预览只画出" in w], [],
                             "\n".join(warns))
            self.assertTrue(os.path.getsize(png) > 10000)

            from PIL import Image
            im = Image.open(png).convert("RGB")
            # getdata() 在 Pillow 14 会移除；新名字是 get_flattened_data()
            reader = getattr(im, "get_flattened_data", None) or im.getdata
            px = set(reader())
            for name, hexv in cs.FLOW_TINTS.items():
                rgb = tuple(int(hexv.lstrip("#")[i:i + 2], 16)
                            for i in (0, 2, 4))
                self.assertIn(rgb, px, "预览里找不到 %s 的填充色 %s"
                                       % (name, hexv))

    def test_rgbcolor_is_tuple_subclass(self):
        """判据自检：确认 `"#%s" % rgb` 真的会炸——否则上一条在空转。"""
        from pptx.dml.color import RGBColor
        c = RGBColor.from_string("AABBCC")
        self.assertIsInstance(c, tuple)
        with self.assertRaises(TypeError):
            "#%s" % c                     # noqa: B018
        self.assertEqual("#" + str(c), "#AABBCC")


class StarterTest(unittest.TestCase):
    """`starter.py` 是给人拷走改的文件，六段必须都能裸跑出来。"""

    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, str(ROOT / "figures"))
        import importlib
        import starter
        cls.mod = importlib.reload(starter)

    def test_all_six_sections_run(self):
        with tempfile.TemporaryDirectory() as d:
            self.mod.OUT = d
            cs.use()
            data = self.mod.load_data()
            for fn in self.mod.STEPS:
                path = fn(data)
                self.assertTrue(os.path.getsize(path) > 5000, path)
            made = [f for f in os.listdir(d) if f.endswith(".png")]
            # 6 张图 + 6 张灰度校样
            self.assertEqual(len(made), 12, sorted(made))

    def test_load_data_supplies_every_key_the_steps_use(self):
        """换数据的人只改 `load_data()`——所以它的键必须齐。缺键要在这里暴露，
        而不是等第 5 段跑到一半 KeyError。"""
        data = self.mod.load_data()
        for key in ("values", "groups", "group_names", "miss", "matrix",
                    "labels", "units", "x", "obs", "fit", "sens", "panel"):
            self.assertIn(key, data)
        self.assertEqual(len(data["labels"]), len(data["units"]),
                         "labels 和 units 必须一一对应——单位跟着变量走")
        self.assertEqual(data["matrix"].shape[1], len(data["labels"]))

    def test_spearman_matches_numpy_on_ranks(self):
        rng = np.random.default_rng(0)
        X = rng.normal(size=(120, 4))
        R = np.apply_along_axis(
            lambda c: np.argsort(np.argsort(c)).astype(float), 0, X)
        np.testing.assert_allclose(self.mod.spearman(X),
                                   np.corrcoef(R, rowvar=False), atol=1e-12)

    def test_spearman_is_invariant_to_monotone_transform(self):
        """选 Spearman 的理由就是这个——单调变换下不变。测出来，别只写在注释里。"""
        rng = np.random.default_rng(1)
        X = rng.random((150, 3)) + 0.1
        a = self.mod.spearman(X)
        b = self.mod.spearman(np.column_stack(
            [np.log(X[:, 0]), X[:, 1] ** 3, np.exp(X[:, 2])]))
        np.testing.assert_allclose(a, b, atol=1e-12)


# **这个块必须留在文件最末尾。** 放在中间的话，`python tests/test_figures.py`
# 执行到那里时后面的 TestCase 还没定义，unittest.main() 收集不到它们：
# **少跑几个类，不报错**。pytest 走 collect 所以看不出来，更难发现。
# （同一个坑在 tests/test_compliance.py 里也踩过一次。）
if __name__ == "__main__":
    unittest.main()
