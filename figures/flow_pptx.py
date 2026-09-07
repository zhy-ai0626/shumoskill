# -*- coding: utf-8 -*-
# docstring 里有 `\includegraphics`，所以整段必须是 raw 字符串。
# 不加 r 前缀，`\i` 是无效转义：Python 3.12 只给 SyntaxWarning，3.15 起报错。
r"""生成**可编辑**的技术路线流程图（.pptx）。

为什么流程图单独走 pptx，而不是和其它图一样用 matplotlib
------------------------------------------------------------
其余 10 张范例都是"数据决定长相"的图——数据一换，图自己就对了，所以代码是
最好的载体。流程图不是：它的内容是**人的判断**（这道题分几步、哪一步反馈到哪一步），
72 小时里跟着模型改三四遍，每次都是挪一个框、改四个字。

这种编辑循环里 pptx 有一条别的格式都没有的优势：**connector 的两端锚在形状上**。
拖动任何一个框，连它的箭头自动跟随。SVG 在 Inkscape 里、canvas 在浏览器里，
挪框之后每一根线都要手动重连。加上 WPS 在参赛队里是默认装备——
pptx 是唯一"队友双击就能改"的格式。

`figures/gallery.py` 里的 `09` 是同一张图的 matplotlib 版本，两边配色一致
（都取 `cumcm_style.FLOW_TINTS`）。用途分工：

- **pptx**：改结构、改文字、给队友改 → 定稿后导出 PDF/PNG 进论文
- **matplotlib**：不想装 Office、要 CI 里可复现、或者要和其它图严格同风格

跑法
----
    python figures/flow_pptx.py                        # 出到 figures/out/
    python figures/flow_pptx.py --output 技术路线.pptx

出来两页：**总体一张 + 每问一张**——44 篇官方展示论文实测，52% 的论文有流程图、
平均 2.3 张，就是这个配比。第 2 页是模板，每个子问复制一份改标题。

导出成图进论文
--------------
WPS / PowerPoint：`文件 → 另存为 → PDF`（矢量，LaTeX 里 `\includegraphics` 直接用），
或 `另存为 → PNG`（选"每张幻灯片"）。**幻灯片尺寸已经设成 20×15 cm**，
不是默认的 16:9——所以导出的图没有大片空白，直接就是插图该有的比例。

已知限制
--------
本文件只**生成**文件，不渲染。生成后请在 WPS / PowerPoint 里打开确认一眼：
中文字体是否正常、连接线有没有跨到框上面。脚本里的 `verify()` 只能查结构
（形状数、连接是否成对、文字是否为空），查不了长相。
"""

from __future__ import annotations

import argparse
import copy
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))

try:
    import _console
except ImportError:                                    # figures/ 被单独拷出去用
    class _console:                                    # noqa: N801
        @staticmethod
        def init() -> None:
            pass

        @staticmethod
        def sym(text: str) -> str:
            return text

import cumcm_style as cs                               # noqa: E402

SLIDE_W_CM = 20.0
SLIDE_H_CM = 15.0

# 中文字体。第一个装了就用第一个——**必须同时写进 latin / ea / cs 三处**，
# 只设 `run.font.name` 只改 latin，中文会掉到 PowerPoint 的默认宋体上，
# 而且**不报错**，你只会觉得"字有点丑"。
CN_FONT = "微软雅黑"
CN_FONT_FALLBACK = "黑体"

# 连接点序号。绝大多数 autoshape 是 0=上 1=左 2=下 3=右。
TOP, LEFT, BOTTOM, RIGHT = 0, 1, 2, 3


# ------------------------------------------------------------------ 底层封装
def _hex(color: str) -> str:
    """`#e3eefb` → `E3EEFB`。pptx 要的是不带井号的大写六位。"""
    return color.lstrip("#").upper()


def _no_shadow(shape) -> None:
    """关掉预设形状自带的投影。

    默认主题给 autoshape 挂了阴影，一张流程图里十个框十片灰影，
    打印出来尤其脏。python-pptx 没有直接 API，改 `shadow.inherit` 即可。
    """
    shape.shadow.inherit = False


def _set_text(shape, lines: list[str], *, size: float = 10.5,
              bold_first: bool = False, color: str = None) -> None:
    """写多行居中文字，并把中文字体钉到 latin/ea/cs 三处。"""
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    # 框里已经很挤，四边内边距压到最小
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)

    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = line
        f = run.font
        f.size = Pt(size if i == 0 else size - 1.0)
        f.bold = bold_first and i == 0
        f.name = CN_FONT
        f.color.rgb = RGBColor.from_string(_hex(color or cs.INK))
        _force_cn_font(run)


def _force_cn_font(run) -> None:
    """把 `<a:ea>` / `<a:cs>` 也设成中文字体。

    `run.font.name` 只写 `<a:latin>`。PowerPoint 对中文字符查的是 `<a:ea>`
    (east asian)，查不到就用主题默认（通常是宋体）——**图上中文会变成宋体
    而不报任何错**。这是 python-pptx 出中文图最常踩的一脚。
    """
    from pptx.oxml.ns import qn

    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", CN_FONT)


def _arrow_head(connector) -> None:
    """给连接线加箭头。

    python-pptx 没有 arrowhead API，得直接往 `<a:ln>` 里塞 `<a:tailEnd>`。
    """
    from pptx.oxml.ns import qn

    ln = connector.line._get_or_add_ln()
    for tag in ("a:tailEnd",):
        el = ln.find(qn(tag))
        if el is None:
            el = ln.makeelement(qn(tag), {})
            ln.append(el)
        el.set("type", "triangle")
        el.set("w", "med")
        el.set("len", "med")


class Deck:
    """薄封装：把"画框 / 连线"变成两行代码，坐标一律用厘米。"""

    def __init__(self):
        from pptx import Presentation
        from pptx.util import Cm

        self.prs = Presentation()
        self.prs.slide_width = Cm(SLIDE_W_CM)
        self.prs.slide_height = Cm(SLIDE_H_CM)
        self._Cm = Cm
        self.slide = None

    def add_slide(self):
        # layout 6 = 空白版式。用带占位符的版式会在图上留下"单击此处添加标题"。
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        return self.slide

    def box(self, x, y, w, h, lines, *, tint="step", size=10.5,
            bold_first=False, rounded=True):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor

        if tint not in cs.FLOW_TINTS:
            raise ValueError("tint 只能是 %s，收到 %r"
                             % ("/".join(cs.FLOW_TINTS), tint))
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            self._Cm(x), self._Cm(y), self._Cm(w), self._Cm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(
            _hex(cs.FLOW_TINTS[tint]))
        shape.line.color.rgb = RGBColor.from_string(_hex(cs.AXIS))
        shape.line.width = self._Cm(0.02)
        _no_shadow(shape)
        if isinstance(lines, str):
            lines = [lines]
        _set_text(shape, lines, size=size, bold_first=bold_first)
        return shape

    def connect(self, src, src_pt, dst, dst_pt, *, elbow=False, dashed=False):
        """连接两个形状。**两端都 connect 上**，这样拖框箭头会跟着走。"""
        from pptx.dml.color import RGBColor
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.enum.shapes import MSO_CONNECTOR

        kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
        # 起止坐标随便给，begin_connect/end_connect 会接管；给个近似值
        # 免得 PowerPoint 在极端情况下路由得很怪。
        c = self.slide.shapes.add_connector(
            kind, src.left, src.top + src.height, dst.left, dst.top)
        c.begin_connect(src, src_pt)
        c.end_connect(dst, dst_pt)
        c.line.color.rgb = RGBColor.from_string(_hex(cs.INK_2))
        c.line.width = self._Cm(0.025)
        if dashed:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        _arrow_head(c)
        return c

    def label(self, x, y, w, h, text, *, size=8.5, align="center",
              color=None):
        """无框文字。用来标边（"解沿用"）、写标题、写图注。"""
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

        box = self.slide.shapes.add_textbox(
            self._Cm(x), self._Cm(y), self._Cm(w), self._Cm(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set_text(box, [text] if isinstance(text, str) else text, size=size,
                  color=color or cs.INK_2)
        for para in tf.paragraphs:
            para.alignment = {"center": PP_ALIGN.CENTER,
                              "left": PP_ALIGN.LEFT}[align]
        return box

    def legend(self, x, y, items):
        """五种角色的色块图例。流程图不用图例框，画小方块最省地方。"""
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor

        step = 3.7
        for k, (tint, text) in enumerate(items):
            sw = self.slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, self._Cm(x + k * step), self._Cm(y),
                self._Cm(0.42), self._Cm(0.30))
            sw.fill.solid()
            sw.fill.fore_color.rgb = RGBColor.from_string(
                _hex(cs.FLOW_TINTS[tint]))
            sw.line.color.rgb = RGBColor.from_string(_hex(cs.AXIS))
            sw.line.width = self._Cm(0.015)
            _no_shadow(sw)
            sw.text_frame.text = ""
            self.label(x + k * step + 0.52, y - 0.12, 2.9, 0.55, text,
                       size=8.0, align="left")

    def save(self, path: str) -> str:
        d = os.path.dirname(path)
        if d:
            os.makedirs(d, exist_ok=True)
        self.prs.save(path)
        return path


LEGEND = (("input", "输入"), ("step", "处理"), ("model", "建模求解"),
          ("check", "校验回检"), ("output", "交付"))


# ------------------------------------------------------------------ 第 1 页
def slide_overall(deck: Deck) -> None:
    """总体技术路线。四条规矩和 matplotlib 版一致：

    1. 层级靠位置和箭头，不靠颜色（灰度稿里五种浅色塌成一片也读得懂）
    2. 连接线两端都 connect 到形状，不要画成自由线段
    3. 反馈边虚线 + 折线，走主干外侧
    4. 框里写"做什么"，不写"用什么"
    """
    deck.add_slide()
    deck.label(0.7, 0.25, 18.6, 0.7,
               "总体技术路线（框 = 做什么；虚线 = 回检不过时的返工路径）",
               size=11.5, align="left", color=cs.INK)
    deck.legend(0.9, 1.25, LEGEND)

    b_in = deck.box(1.0, 2.0, 18.0, 1.15, "题面 + 附件数据", tint="input")
    b_sc = deck.box(1.0, 3.75, 18.0, 1.55,
                    ["附件结构体检 → 判题型",
                     "成分？重复测量？数据口径？"], tint="step")

    qw, qgap = 5.6, 0.7
    b_q1 = deck.box(1.0, 6.3, qw, 1.85,
                    ["Q1 基准模型", "最简情形先跑通"], tint="model",
                    bold_first=True)
    b_q2 = deck.box(1.0 + qw + qgap, 6.3, qw, 1.85,
                    ["Q2 加不确定性", "随机 / 鲁棒"], tint="model",
                    bold_first=True)
    b_q3 = deck.box(1.0 + 2 * (qw + qgap), 6.3, qw, 1.85,
                    ["Q3 加耦合", "变量间相关性"], tint="model",
                    bold_first=True)

    b_ck = deck.box(3.4, 9.3, 13.2, 1.6,
                    ["跨问一致性回检",
                     "把「上界 / 不可行」类断言用最终解代回"], tint="check")
    b_se = deck.box(3.4, 11.5, 13.2, 1.15,
                    "灵敏度 + 鲁棒性（按影响幅度排序）", tint="step")
    b_out = deck.box(3.4, 13.2, 13.2, 1.15,
                     "结论、决策建议与交付文件", tint="output")

    deck.connect(b_in, BOTTOM, b_sc, TOP)
    # b_sc 横跨整行，它的下边中点不在 Q1 上方 → 用折线连接器，PowerPoint 会
    # 自己走直角，比斜线好读，而且拖动 Q1 时会重新路由。
    deck.connect(b_sc, BOTTOM, b_q1, TOP, elbow=True)
    deck.connect(b_q1, RIGHT, b_q2, LEFT)
    deck.connect(b_q2, RIGHT, b_q3, LEFT)
    deck.label(1.0 + qw - 0.35, 5.75, 2.0, 0.5, "解沿用", size=8.0)
    deck.label(1.0 + 2 * qw + qgap - 0.35, 5.75, 2.0, 0.5, "解沿用", size=8.0)
    for b in (b_q1, b_q2, b_q3):
        deck.connect(b, BOTTOM, b_ck, TOP)
    deck.connect(b_ck, BOTTOM, b_se, TOP)
    deck.connect(b_se, BOTTOM, b_out, TOP)
    # 反馈环：回检不过回到建模。虚线折线，走左侧空白，不跨任何主干线。
    deck.connect(b_ck, LEFT, b_q1, BOTTOM, elbow=True, dashed=True)
    deck.label(0.35, 8.55, 3.1, 0.55, "不一致 → 回到建模", size=8.0)

    deck.label(0.7, 14.5, 18.6, 0.45,
               "画法依据：44 篇官方展示论文中 52% 有流程图，"
               "平均 2.3 张（总体一张 + 每问一张）",
               size=7.5, align="left")


# ------------------------------------------------------------------ 第 2 页
def slide_per_question(deck: Deck) -> None:
    """单个子问的求解流程模板。**每个子问复制这一页，改标题和框里的内容。**

    为什么值得单独一页：官方讲评连年点名"没有给出具体算法""计算方法不清楚"
    （2022A 评阅问题 9 条里 5 条是呈现类）。把"算法 + 参数""自检项"
    画成显式的框，写论文时就不会漏——**图上空着的框，就是论文里缺的一段**。
    """
    deck.add_slide()
    deck.label(0.7, 0.25, 18.6, 0.7,
               "Qi 求解流程（每问复制一页，改标题与框内文字）",
               size=11.5, align="left", color=cs.INK)
    deck.legend(0.9, 1.25, LEGEND)

    cx, cw = 4.2, 9.4
    rows = [
        ("input", ["输入", "上游结果 + 附件第 X 列 + 题给参数"]),
        ("model", ["数学模型", "决策变量 / 目标函数 / 约束（符号级）"]),
        # 用 model 而不是 step：图例里 model 那一档写的是"建模求解"，
        # 求解算法归到"处理"会和自己的图例打架。
        ("model", ["求解算法", "算法名 + 关键参数 + 收敛判据"]),
        ("check", ["自检", "恒等式 / 量纲 / 极限情形 / 逐自由度覆盖"]),
        ("output", ["结果与分析", "表 + 图 + 一段「这说明什么」的分析"]),
    ]
    y, gap, bh = 2.1, 0.75, 1.75
    boxes = []
    for tint, lines in rows:
        boxes.append(deck.box(cx, y, cw, bh, lines, tint=tint,
                              bold_first=True))
        y += bh + gap

    for a, b in zip(boxes, boxes[1:]):
        deck.connect(a, BOTTOM, b, TOP)
    # 自检不过 → 回到模型。虚线走右侧。
    deck.connect(boxes[3], RIGHT, boxes[1], RIGHT, elbow=True, dashed=True)
    deck.label(14.0, 5.9, 5.3, 1.1,
               ["自检不过", "→ 改模型，不是改数"], size=8.0)

    # 左侧提示条：这三句是讲评里连年重复的失分点，画在图边当检查表
    deck.label(0.3, 3.0, 3.6, 4.6,
               ["写论文时对着这张图查：", "",
                "· 算法框空着 → 论文缺「具体算法」",
                "· 自检框空着 → 缺自检，或写成了「局限」",
                "· 分析框只有表 → 缺「这说明什么」的那段"],
               size=8.0, align="left")

    deck.label(0.7, 14.5, 18.6, 0.45,
               "官方连年点名：计算方法不清楚 / 没有给出具体算法 / "
               "没有对结果进行分析（2022A 评阅问题 9 条里 5 条是呈现类）",
               size=7.5, align="left")


# ------------------------------------------------------------------ 校验
def verify(path: str) -> list[str]:
    """结构性自检。**查不了长相**——那要你在 WPS 里开一眼。

    查三件在生成阶段最容易静默出错的事：
    1. 连接线两端是否都真的绑到了形状（只 add_connector 不 connect 的话，
       线画出来一模一样，但**拖框不跟随**——这正是选 pptx 的唯一理由）；
    2. 有没有空文字框（`_set_text` 传空 list 时会静默留个空框）；
    3. 中文 run 是否都设了 `<a:ea>`（没设不报错，只是中文变宋体）。
    """
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(path)
    problems: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        n_box = n_conn = n_bound = 0
        for shape in slide.shapes:
            if shape.shape_type is not None and shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if not txt and shape.width > 0 and shape.name.startswith(
                        ("TextBox", "Rounded")):
                    problems.append("第 %d 页有空文字框 %s" % (i, shape.name))
            if shape.element.tag.endswith("}cxnSp"):
                n_conn += 1
                nv = shape.element.find(
                    qn("p:nvCxnSpPr") + "/" + qn("p:cNvCxnSpPr"))
                has_st = nv is not None and nv.find(qn("a:stCxn")) is not None
                has_end = nv is not None and nv.find(qn("a:endCxn")) is not None
                if has_st and has_end:
                    n_bound += 1
                else:
                    problems.append(
                        "第 %d 页有一根连接线没绑住两端（拖框不会跟随）" % i)
            elif shape.has_text_frame:
                n_box += 1
            for para in getattr(shape, "text_frame", None).paragraphs \
                    if shape.has_text_frame else []:
                for run in para.runs:
                    if any("一" <= ch <= "鿿" for ch in run.text):
                        rPr = run._r.find(qn("a:rPr"))
                        if rPr is None or rPr.find(qn("a:ea")) is None:
                            problems.append(
                                "第 %d 页有中文 run 没设 ea 字体：%r"
                                % (i, run.text[:12]))
        print("  第 %d 页：%d 个文本形状，%d 根连接线（%d 根两端已绑定）"
              % (i, n_box, n_conn, n_bound))
    return problems


PT_CM = 2.54 / 72.0            # 1 pt = 0.0353 cm


def _text_cm(text: str, pt: float) -> float:
    """量一行字在 `pt` 字号下的宽度（厘米）。

    用 matplotlib 的真实字体度量，不是"按字数 × 系数"估——中英混排时
    估算能差 40%，而这个数是判"文字会不会挤出框"的唯一依据。
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=(1, 1), dpi=100)
    t = fig.text(0, 0, text, fontsize=pt,
                 fontname=cs.available_cn_font() or "DejaVu Sans")
    fig.canvas.draw()
    bb = t.get_window_extent(fig.canvas.get_renderer())
    plt.close(fig)
    return bb.width / 100.0 * 2.54      # px @100dpi → inch → cm


def preview(pptx_path: str, png_path: str) -> list[str]:
    r"""把 pptx 里的**真实坐标**读回来渲染一张预览图，并量文字是否溢出。

    为什么需要这一步：本机没有 Office/WPS/LibreOffice，`.pptx` 生成完是个
    黑盒——框有没有压到一起、字有没有挤出框，`verify()` 一概查不出来。
    这里从文件里读回每个形状的 left/top/width/height 和字号，用 matplotlib
    按 1:1 比例重画，就把"看不见"变成"只有 PowerPoint 的折线路由看不见"。

    预览**不等于**真实渲染：PowerPoint 的 elbow 连接器自己走直角，这里一律
    画直线；换行算法也不完全一致。它能查的是位置、尺寸、文字长度这三件事。
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL
    from pptx.oxml.ns import qn

    cs.use()
    prs = Presentation(pptx_path)
    W = prs.slide_width / 360000.0          # EMU → cm
    H = prs.slide_height / 360000.0
    slides = list(prs.slides)
    fig, axes = plt.subplots(1, len(slides),
                             figsize=(W / 2.54 * len(slides), H / 2.54))
    if len(slides) == 1:
        axes = [axes]

    warnings: list[str] = []
    n_filled = 0
    for ax, slide in zip(axes, slides):
        ax.set_xlim(0, W)
        ax.set_ylim(H, 0)                   # pptx 的 y 轴向下
        ax.set_aspect("equal")
        ax.set_axis_off()
        ax.grid(False)
        ax.add_patch(plt.Rectangle((0, 0), W, H, facecolor=cs.SURFACE,
                                   edgecolor=cs.AXIS, linewidth=0.6))

        by_id = {}
        for shape in slide.shapes:
            if shape.element.tag.endswith("}cxnSp"):
                continue
            x = shape.left / 360000.0
            y = shape.top / 360000.0
            w = shape.width / 360000.0
            h = shape.height / 360000.0
            by_id[shape.shape_id] = (x, y, w, h)
            # **按填充类型判断，不按形状名判断**，也不用 try/except 兜。
            # 第一版写的是 `"#%s" % shape.fill.fore_color.rgb` 并用
            # `except Exception` 兜底：`RGBColor` 是 **tuple 子类**，
            # `"#%s" % rgb` 会把三个分量当成三个参数塞进一个占位符而抛
            # TypeError，于是**每个框都静默变成画布白** ——
            # 预览图看着"都是白框"，pptx 本身其实是对的。
            # 教训：兜底把 bug 变成了错的输出，而不是变成一条报错。
            rgb = None
            if shape.fill.type == MSO_FILL.SOLID:
                rgb = "#" + str(shape.fill.fore_color.rgb)
            filled = rgb is not None
            if filled:
                n_filled += 1
                ax.add_patch(plt.Rectangle((x, y), w, h, facecolor=rgb,
                                           edgecolor=cs.AXIS, linewidth=0.7,
                                           zorder=2))
            if not shape.has_text_frame:
                continue
            lines = [p.text for p in shape.text_frame.paragraphs if p.text]
            if not lines:
                continue
            sizes = [(p.runs[0].font.size.pt if p.runs and p.runs[0].font.size
                      else 10.5) for p in shape.text_frame.paragraphs if p.text]
            # 量溢出：任何一行的实际宽度超过框宽（扣掉 0.12 cm 内边距）就报
            usable = max(w - 0.12, 0.1)
            need_rows = 0
            for line, pt in zip(lines, sizes):
                tw = _text_cm(line, pt)
                rows = max(1, int(tw / usable) + (1 if tw % usable else 0))
                need_rows += rows
                if filled and tw > usable * 2.2:
                    warnings.append(
                        "文字过长，%s 里 %r 需要 %.1f cm 而框只有 %.1f cm"
                        % (shape.name, line[:16], tw, usable))
            line_cm = max(sizes) * PT_CM * 1.5
            if filled and need_rows * line_cm > h + 0.05:
                warnings.append(
                    "文字换行后高度溢出：%s 需要约 %.2f cm，框高 %.2f cm"
                    % (shape.name, need_rows * line_cm, h))
            for k, (line, pt) in enumerate(zip(lines, sizes)):
                yy = y + h / 2 + (k - (len(lines) - 1) / 2) * pt * PT_CM * 1.5
                ax.text(x + w / 2 if filled else x + w / 2, yy, line,
                        ha="center", va="center", fontsize=pt * 0.92,
                        color=cs.INK if filled else cs.INK_2, zorder=4)

        # 连接线：按绑定的连接点画直线（PowerPoint 的折线路由这里画不出来）
        for shape in slide.shapes:
            if not shape.element.tag.endswith("}cxnSp"):
                continue
            nv = shape.element.find(qn("p:nvCxnSpPr") + "/"
                                    + qn("p:cNvCxnSpPr"))
            st, en = nv.find(qn("a:stCxn")), nv.find(qn("a:endCxn"))
            if st is None or en is None:
                continue

            def pt_of(node):
                geo = by_id.get(int(node.get("id")))
                if geo is None:
                    return None
                x, y, w, h = geo
                return {0: (x + w / 2, y), 1: (x, y + h / 2),
                        2: (x + w / 2, y + h), 3: (x + w, y + h / 2)}[
                            int(node.get("idx"))]

            p, q = pt_of(st), pt_of(en)
            if p is None or q is None:
                continue
            # dash_style 是枚举成员而不是字符串，`.lower()` 会 AttributeError
            dashed = "DASH" in str(shape.line.dash_style or "").upper()
            ax.annotate("", xy=q, xytext=p, zorder=3,
                        arrowprops=dict(arrowstyle="-|>", color=cs.INK_2,
                                        linewidth=0.8, shrinkA=1, shrinkB=1,
                                        linestyle="--" if dashed else "-",
                                        mutation_scale=8))

    # **预览自己也要有自检。** "所有框都是白的"是这个函数踩过的真实 bug
    # （见上面那段注释），而白框和"本来就没填色"在图上一模一样。
    # 所以显式对账：pptx 里有几个实心形状，预览就该画出几个色块。
    n_solid = sum(1 for s in slides for sh in s.shapes
                  if not sh.element.tag.endswith("}cxnSp")
                  and sh.fill.type == MSO_FILL.SOLID)
    if n_solid and n_filled < n_solid:
        warnings.append(
            "预览只画出 %d / %d 个实心形状的填充色——预览渲染有问题，"
            "不是 pptx 有问题" % (n_filled, n_solid))

    cs.add_units_note(fig, "由 .pptx 的真实坐标重绘的预览。"
                           "PowerPoint 的折线连接器在这里一律画成直线，"
                           "长相以 WPS 里打开为准。")
    fig.tight_layout()
    cs.save(fig, png_path, proof=False)
    return warnings


def main(argv: list[str] | None = None) -> int:
    _console.init()
    ap = argparse.ArgumentParser(
        description="生成可编辑的技术路线流程图 (.pptx)")
    default = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "out", "技术路线_可编辑.pptx")
    ap.add_argument("--output", default=default, help="输出 .pptx 路径")
    ap.add_argument("--preview", nargs="?", const="", default=None,
                    metavar="PNG",
                    help="额外渲染一张预览 PNG（从 pptx 的真实坐标重绘），"
                         "并检查文字有没有挤出框。不给路径就放在 pptx 旁边")
    args = ap.parse_args(argv)

    try:
        import pptx  # noqa: F401
    except ImportError:
        print("需要 python-pptx：python -m pip install python-pptx",
              file=sys.stderr)
        return 2

    deck = Deck()
    slide_overall(deck)
    slide_per_question(deck)
    path = deck.save(args.output)
    print("已生成 %s（%.1f×%.1f cm，2 页）" % (path, SLIDE_W_CM, SLIDE_H_CM))

    problems = verify(path)

    if args.preview is not None:
        png = args.preview or (os.path.splitext(path)[0] + "_预览.png")
        overflow = preview(path, png)
        print("预览图：%s" % png)
        if overflow:
            print(_console.sym("  ⚠ 文字排布有 %d 处可疑：" % len(overflow)))
            for w in overflow[:8]:
                print("      " + w)
            problems.extend(overflow)
        else:
            print(_console.sym("  ✓ 每个框的文字都装得下。"))

    if problems:
        print(_console.sym("\n✗ 自检有问题："))
        for p in problems[:12]:
            print("    " + p)
        return 1
    print(_console.sym("\n✓ 自检通过：连接线两端都绑在形状上，"
                       "拖动框时箭头会跟随。"))
    print("下一步：在 WPS / PowerPoint 里打开确认长相（脚本查不了折线路由），"
          "改完 另存为 → PDF 进论文。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
